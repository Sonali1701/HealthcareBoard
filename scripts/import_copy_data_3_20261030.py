"""Audit and import Copy of Data 3.xlsx into the production provider directory.

Dry-run is the default. The import is deterministic and rerunnable. Existing
email or phone identities are never updated; invalid and conflicting rows are
excluded before the database transaction starts.
"""
from __future__ import annotations

import argparse
import hashlib
import re
import uuid
from collections import Counter, defaultdict
from pathlib import Path

import psycopg
from openpyxl import load_workbook
from sqlalchemy.engine import make_url

from app.importers.parsing import is_real_name


EXPECTED_HOST = "ep-nameless-flower-ah7umr3s-pooler.c-3.us-east-1.aws.neon.tech"
SOURCE = "xlsx_copy_data_3_20261030"
NAMESPACE = uuid.UUID("1fc94468-dd68-45bb-94fc-69ef93c0733c")
HEADERS = ("Name", "Number", "City", "State", "Speciality", "Email")
EMAIL = re.compile(r"^[^@\s]+@[^@\s]+\.[^@\s]+$")
STATES = set(
    "AL AK AZ AR CA CO CT DE DC FL GA HI ID IL IN IA KS KY LA ME MD MA MI MN "
    "MS MO MT NE NV NH NJ NM NY NC ND OH OK OR PA RI SC SD TN TX UT VT VA WA "
    "WV WI WY".split()
)

NONCLINICAL = {
    "other", "clerk", "admin assistant", "medical coder", "supervisor",
    "healthcare administrator", "charge/supervisor/manager",
    "director of nursing - don", "asst. director of nursing - adon",
}
NURSING = {
    "ms", "ms tele", "icu", "psych", "er", "pcu", "or",
    "correctional health", "tele", "l&d", "pacu", "ltc", "dialysis",
    "mother baby", "micu", "rehab", "picu", "oncology", "case manager",
    "cvicu", "clinic", "endoscopy", "home health", "peds", "ambulatory",
    "ir", "cvor", "charge", "ccu", "pre/post/peri op care", "acute care",
    "sicu", "snf", "ms ortho", "infusion", "neuro icu", "hospice",
    "postpartum", "peds er", "wound care", "peds ms", "ltac", "float pool",
    "infection control", "apheresis", "rn general", "nurse educator",
    "neuro ms", "assisted living",
}
ALLIED_TECH = {
    "rrt", "crt", "certified surgical/scrub tech", "sterile processing tech",
    "certified sterile processing tech", "phlebotomist", "mt", "mlt",
    "ct tech", "rad tech", "cath lab", "medical assistant", "cma",
    "x-ray tech", "certified dialysis tech", "mri tech",
    "surgical/scrub tech (non certified)", "echo tech", "dialysis tech",
    "anesthesia tech", "radiology", "monitor/telemetry tech", "pct",
    "nuclear med tech", "eeg tech", "cath lab tech", "histo tech",
    "endo tech", "lab assistant/coordinator", "medical physicist",
    "cardiac sonographer", "mammo tech", "anesthesia", "ir tech", "csfa",
    "vascular tech", "patient care associate", "ekg tech",
    "sterile processing manager", "fa", "surgical first assistant",
}


def clean(value: object) -> str:
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return " ".join(str(value or "").split())


def parse_name(value: str) -> tuple[str, str] | None:
    # Credentials follow a comma in this workbook ("First Last, RN BSN").
    name = clean(value).split(",", 1)[0].strip()
    parts = [part.strip(" .") for part in name.split() if part.strip(" .")]
    if len(parts) < 2:
        return None
    first, last = parts[0], parts[-1]
    if len(first) > 100 or len(last) > 100 or not is_real_name(first, last):
        return None
    return first, last


def taxonomy(specialty: str) -> tuple[str, str, str] | None:
    """Return profession type, job category, and broad provider category."""
    key = clean(specialty).casefold()
    if not key or key in NONCLINICAL:
        return None
    if key.startswith("np "):
        return "NP", "Nursing", "APP"
    if key == "lpn/lvn general":
        return "LPN", "Nursing", "Nursing"
    if key in {"cna", "nurse aide"}:
        return "CNA", "Nursing", "Nursing"
    if key in NURSING or key in {"nicu"}:
        return "RN", "Nursing", "Nursing"
    if key in {"pharmacy tech", "certified pharmacy tech"}:
        return "Pharmacy Technician", "Pharmacy", "Allied"
    if key == "pharmacist":
        return "PharmD", "Pharmacy", "Allied"
    if key in {"lcsw", "msw", "social worker", "social worker - case management"}:
        return "Social Worker", "Social Work & Community Services", "Allied"
    if key == "behavioral health tech":
        return "Behavioral Health Technician", "Mental & Behavioural Health", "Allied"
    if key in {"ot", "pt", "cota", "pta", "radiation therapist"}:
        profession = {
            "ot": "OT", "pt": "PT", "cota": "COTA", "pta": "PTA",
            "radiation therapist": "Radiation Therapist",
        }[key]
        return profession, "Therapy", "Allied"
    if key in {"registered dietician", "dietician"}:
        return "Dietitian", "Allied Health & Technicians", "Allied"
    if key == "paramedic/emt":
        return "Paramedic", "Medical Transportation", "Allied"
    if key == "sonographer/ultrasound tech":
        return "Sonographer", "Allied Health & Technicians", "Allied"
    if key in ALLIED_TECH:
        return clean(specialty)[:50], "Allied Health & Technicians", "Allied"
    return None


def load(path: Path) -> tuple[list[tuple], Counter]:
    book = load_workbook(path, read_only=True, data_only=True)
    if len(book.worksheets) != 1:
        raise ValueError(f"Expected one worksheet, found {len(book.worksheets)}")
    rows = book.active.iter_rows(values_only=True)
    header = [clean(value) for value in next(rows)]
    if tuple(header[:len(HEADERS)]) != HEADERS or any(header[len(HEADERS):]):
        raise ValueError(f"Unexpected columns: {header}")
    index = {name: header.index(name) for name in HEADERS}
    grouped: dict[str, list[dict]] = defaultdict(list)
    stats: Counter = Counter()

    for row_number, row in enumerate(rows, 2):
        if not any(value not in (None, "") for value in row):
            continue
        stats["workbook_rows"] += 1
        def value(name: str) -> str:
            return clean(row[index[name]])
        email = value("Email").casefold()
        parsed_name = parse_name(value("Name"))
        phone_digits = re.sub(r"\D", "", value("Number"))[-10:]
        state = value("State").upper()
        specialty = value("Speciality")
        resolved = taxonomy(specialty)
        if not EMAIL.fullmatch(email) or len(email) > 255:
            stats["invalid_email"] += 1
            continue
        if not parsed_name:
            stats["invalid_name"] += 1
            continue
        if state not in STATES:
            stats["invalid_or_non_us_state"] += 1
            continue
        if not resolved:
            stats["unmapped_or_nonclinical_specialty"] += 1
            continue
        if phone_digits and len(phone_digits) != 10:
            stats["invalid_phone"] += 1
            continue
        first, last = parsed_name
        profession, category, provider_category = resolved
        grouped[email].append({
            "row": row_number, "first": first, "last": last,
            "phone": phone_digits, "city": value("City")[:120], "state": state,
            "specialty": specialty[:100], "profession": profession[:50],
            "category": category, "provider_category": provider_category,
        })
    book.close()

    candidates: list[dict] = []
    for email, copies in grouped.items():
        identities = {
            (item["first"].casefold(), item["last"].casefold(), item["phone"],
             item["state"], item["specialty"].casefold())
            for item in copies
        }
        if len(identities) != 1:
            stats["conflicting_repeated_email"] += 1
            continue
        stats["repeated_rows_removed"] += len(copies) - 1
        item = min(copies, key=lambda candidate: candidate["row"])
        item["email"] = email
        candidates.append(item)

    by_phone: dict[str, list[dict]] = defaultdict(list)
    for item in candidates:
        if item["phone"]:
            by_phone[item["phone"]].append(item)
    rejected: set[str] = set()
    for matches in by_phone.values():
        if len(matches) < 2:
            continue
        names = {(item["first"].casefold(), item["last"].casefold()) for item in matches}
        if len(names) > 1:
            rejected.update(item["email"] for item in matches)
            stats["shared_phone_identity_conflict"] += len(matches)
        else:
            keep = min(matches, key=lambda item: item["row"])
            rejected.update(item["email"] for item in matches if item is not keep)
            stats["same_identity_multiple_emails"] += len(matches) - 1

    result = []
    for item in candidates:
        if item["email"] in rejected:
            continue
        profile_id = str(uuid.uuid5(NAMESPACE, item["email"]))
        phone = item["phone"]
        formatted_phone = (
            f"({phone[:3]}) {phone[3:6]}-{phone[6:]}" if phone else None
        )
        result.append((
            profile_id, item["row"], item["email"], item["first"], item["last"],
            formatted_phone, item["phone"] or None, item["city"] or None,
            item["state"], item["specialty"], item["profession"],
            item["category"], item["provider_category"],
        ))
    stats["valid_unique_rows"] = len(result)
    return result, stats


def database_url(env_file: Path) -> str:
    values = {}
    for raw in env_file.read_text(encoding="utf-8-sig").splitlines():
        line = raw.strip()
        if not line or line.startswith("#") or "=" not in line:
            continue
        key, value = line.split("=", 1)
        values[key.strip()] = value.strip().strip('"').strip("'")
    return values["DATABASE_URL"]


CREATE_STAGE = """
CREATE TEMP TABLE copy_data_3_stage (
 profile_id varchar(36) PRIMARY KEY, source_row integer NOT NULL,
 email varchar(255) NOT NULL UNIQUE, first_name varchar(100) NOT NULL,
 last_name varchar(100) NOT NULL, phone varchar(30), phone_digits varchar(10),
 city varchar(120), state_code varchar(2) NOT NULL, specialty varchar(100) NOT NULL,
 profession_type varchar(50) NOT NULL, job_category varchar(255) NOT NULL,
 provider_category varchar(20) NOT NULL
) ON COMMIT DROP
"""

CREATE_AUDIT = """
CREATE TABLE IF NOT EXISTS profile_copy_data_3_import_20261030 (
 profile_id varchar(36) PRIMARY KEY REFERENCES profiles(profile_id) ON DELETE CASCADE,
 source_file text NOT NULL, source_sha256 varchar(64) NOT NULL,
 source_row integer NOT NULL, imported_at timestamptz NOT NULL DEFAULT now()
)
"""

CREATE_REVIEW = """
CREATE TABLE IF NOT EXISTS profile_copy_data_3_review_20261030 (
 email varchar(255) PRIMARY KEY, source_row integer NOT NULL,
 matched_profile_id varchar(36), reason varchar(40) NOT NULL,
 recorded_at timestamptz NOT NULL DEFAULT now()
)
"""

MATCHES = """
WITH email_matches AS (
 SELECT s.profile_id, min(p.profile_id) AS email_id
 FROM profiles p JOIN copy_data_3_stage s ON lower(btrim(p.email)) = s.email
 GROUP BY s.profile_id
), phone_matches AS (
 SELECT s.profile_id, min(p.profile_id) AS phone_id
 FROM profiles p JOIN copy_data_3_stage s
   ON s.phone_digits IS NOT NULL
  AND right(regexp_replace(coalesce(p.phone,''),'[^0-9]','','g'),10) = s.phone_digits
 GROUP BY s.profile_id
), matched AS (
 SELECT s.*, e.email_id, ph.phone_id
 FROM copy_data_3_stage s
 LEFT JOIN email_matches e USING (profile_id)
 LEFT JOIN phone_matches ph USING (profile_id)
)
"""

INSERT = MATCHES + """
, eligible AS (
 SELECT m.* FROM matched m
 WHERE email_id IS NULL AND phone_id IS NULL
)
INSERT INTO profiles (
 profile_id, first_name, last_name, headline, phone, email, specialty,
 profession_type, provider_category, job_category, is_listable, screen_reason,
 screened_at, city, state_code, lat, lng, years_experience, open_to_work,
 job_type_prefs, completion_score, source, capture_source, search_text,
 created_at, updated_at
)
SELECT e.profile_id, e.first_name, e.last_name,
       concat(e.specialty, ' · ', e.profession_type), e.phone, e.email,
       e.specialty, e.profession_type, e.provider_category, e.job_category,
       true, 'structured_contact_unverified', now(), e.city, e.state_code,
       geo.lat, geo.lng, 0, false, '[]'::json, 55,
       'json_import'::profilesource, %s,
       lower(concat_ws(' ', e.first_name, e.last_name, e.specialty,
                       e.profession_type, e.job_category, e.provider_category,
                       e.city, e.state_code)), now(), now()
FROM eligible e
LEFT JOIN LATERAL (
 SELECT c.lat, c.lng FROM city_centroids c
 WHERE c.city_lower = lower(btrim(e.city)) AND c.state_code = e.state_code
 LIMIT 1
) geo ON true
ON CONFLICT DO NOTHING
RETURNING profile_id
"""


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", type=Path)
    parser.add_argument("--env-file", type=Path, required=True)
    parser.add_argument("--apply", action="store_true")
    args = parser.parse_args()
    if not args.path.is_file() or not args.env_file.is_file():
        raise FileNotFoundError(args.path if not args.path.is_file() else args.env_file)
    rows, stats = load(args.path)
    print("WORKBOOK_AUDIT", dict(stats), flush=True)

    url = make_url(database_url(args.env_file))
    if url.host != EXPECTED_HOST or url.database != "neondb":
        raise RuntimeError("Refusing unexpected database target")
    dsn = url.set(drivername="postgresql").render_as_string(hide_password=False)
    digest = hashlib.sha256(args.path.read_bytes()).hexdigest()
    with psycopg.connect(dsn, connect_timeout=10) as conn:
        with conn.cursor() as cur:
            cur.execute("SET LOCAL statement_timeout = '900s'")
            cur.execute(CREATE_STAGE)
            cur.execute("CREATE INDEX copy_data_3_stage_phone ON copy_data_3_stage(phone_digits)")
            with cur.copy("COPY copy_data_3_stage FROM STDIN") as copy:
                for row in rows:
                    copy.write_row(row)
            cur.execute(MATCHES + """
                SELECT count(*) FILTER (WHERE email_id IS NOT NULL),
                       count(*) FILTER (WHERE phone_id IS NOT NULL),
                       count(*) FILTER (WHERE email_id IS NOT NULL AND phone_id IS NOT NULL
                                         AND email_id <> phone_id),
                       count(*) FILTER (WHERE email_id IS NULL AND phone_id IS NULL)
                FROM matched
            """)
            email_matches, phone_matches, identity_conflicts, eligible = cur.fetchone()
            print("DATABASE_AUDIT", {
                "existing_email": email_matches, "existing_phone": phone_matches,
                "email_phone_identity_conflicts": identity_conflicts,
                "eligible_new_profiles": eligible,
            }, flush=True)
            if not args.apply:
                conn.rollback()
                print("DRY_RUN_ONLY", flush=True)
                return

            cur.execute(CREATE_AUDIT)
            cur.execute(CREATE_REVIEW)
            cur.execute(MATCHES + """
                INSERT INTO profile_copy_data_3_review_20261030
                    (email, source_row, matched_profile_id, reason)
                SELECT email, source_row, coalesce(email_id, phone_id),
                       CASE WHEN email_id IS NOT NULL AND phone_id IS NOT NULL
                                      AND email_id <> phone_id THEN 'email_phone_conflict'
                            WHEN email_id IS NOT NULL THEN 'existing_email'
                            ELSE 'existing_phone' END
                FROM matched WHERE email_id IS NOT NULL OR phone_id IS NOT NULL
                ON CONFLICT (email) DO UPDATE SET
                    source_row=excluded.source_row,
                    matched_profile_id=excluded.matched_profile_id,
                    reason=excluded.reason,
                    recorded_at=now()
            """)
            review_rows = cur.rowcount
            cur.execute(INSERT, (SOURCE,))
            inserted_ids = [row[0] for row in cur.fetchall()]
            cur.execute("""
                INSERT INTO profile_copy_data_3_import_20261030
                    (profile_id, source_file, source_sha256, source_row)
                SELECT p.profile_id, %s, %s, s.source_row
                FROM copy_data_3_stage s JOIN profiles p ON p.profile_id=s.profile_id
                ON CONFLICT (profile_id) DO NOTHING
            """, (args.path.name, digest))
            audit_rows = cur.rowcount
            cur.execute("""
                SELECT count(*), count(*) FILTER (WHERE p.is_listable),
                       count(*) FILTER (WHERE p.lat IS NOT NULL AND p.lng IS NOT NULL),
                       count(*) FILTER (WHERE p.email IS NOT NULL),
                       count(*) FILTER (WHERE p.phone IS NOT NULL)
                FROM profile_copy_data_3_import_20261030 a
                JOIN profiles p USING (profile_id)
            """)
            verified = cur.fetchone()
            print("APPLIED", {
                "profiles_inserted": len(inserted_ids),
                "review_rows_recorded": review_rows,
                "audit_rows_inserted": audit_rows,
                "verified_total_visible_geocoded_email_phone": verified,
            }, flush=True)


if __name__ == "__main__":
    main()
